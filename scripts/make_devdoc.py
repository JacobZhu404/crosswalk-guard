#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《开发过程文档》Word 版 —— 既讲系统怎么运作, 也讲这套开发方法论怎么工作。
专业表述 + 通俗解释并行。
用法: python scripts/make_devdoc.py [--out docs/manuals]
"""
import os
import argparse
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def _set_base_font(doc, font="微软雅黑", size=11):
    style = doc.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(size)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), font)


def _title(doc, text, subtitle=None):
    h = doc.add_heading(text, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def _p(doc, text, bold=False, size=None, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    return p


def _plain(doc, text):
    """通俗解释段: 浅灰斜体, 和专业表述区分。"""
    p = doc.add_paragraph()
    r = p.add_run("通俗讲：" + text)
    r.italic = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x55, 0x66, 0x88)
    p.paragraph_format.left_indent = Inches(0.3)
    return p


def _bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    return p


def _num(doc, text):
    return doc.add_paragraph(text, style="List Number")


def _table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(10)
    return t


def build(path):
    doc = Document()
    _set_base_font(doc)

    _title(doc, "斑马线违章检测系统 · 开发过程文档",
           "系统如何运作 + 我们如何开发 · 系统版本 2.3.0 · 生成于 2026-09-30")

    # ---- 引言 ----
    doc.add_heading("写在前面", level=1)
    _p(doc, "这份文档有两条线。第一条讲“系统是怎么运作的”——它拿到一段视频后，如何一步步判断出违章。"
            "第二条，也是本文的重点，讲“我们是怎么把它开发出来的”——不是靠一次性写好代码，"
            "而是靠一套可以反复迭代、每一步都能验证的工程方法。")
    _p(doc, "每个专业段落后面，都配一句“通俗讲”，方便不同读者各取所需。")

    # ======================================================
    # 第一部分: 系统怎么运作(简版)
    # ======================================================
    doc.add_heading("第一部分 · 系统是怎么运作的", level=1)

    _p(doc, "系统把“判断违章”这件复杂的事，拆成一条可解释的流水线，逐帧看视频、最后综合判断：", bold=True)
    _table(doc,
           ["步骤", "做什么", "怎么做（技术）"],
           [
               ["① 找车", "逐帧框出车辆并持续跟踪", "YOLOv8n 检测 + IoU 跟踪，每辆车一个 track_id"],
               ["② 找斑马线", "识别斑马线区域", "CrosswalkDetectorV2，四边形探针 + 时序聚合"],
               ["③ 看灯", "判断行人灯颜色", "HSV 亮斑颜色法 + 逐视频位置先验 + 时序平滑"],
               ["④ 认牌", "读涉事车车牌", "HyperLPR3 + 多帧加权投票"],
               ["⑤ 判定", "三条件同时成立才算违章", "行人绿灯 ∧ 车静止 ∧ 压斑马线，且持续足够时长"],
           ])
    _plain(doc, "系统像一位交通协管，先看清“有哪些车、斑马线在哪、灯什么颜色、车牌是多少”，"
                "再按一条铁规矩下结论：只有行人绿灯时、车还停在斑马线上不走，才算违章。")

    _p(doc, "两个关键的“谨慎”设计：", bold=True)
    _bullet(doc, "没拍到红绿灯时，绝不瞎判违章，而是标“待复核”，交给人看。")
    _bullet(doc, "车牌读不清时，宁可留空，也不乱填一个号码去冤枉车主。")
    _plain(doc, "它的原则是“宁可少报，不可错报”——尤其绝不冤枉守规矩的车。")

    # ======================================================
    # 第二部分: 开发方法论(重点)
    # ======================================================
    doc.add_heading("第二部分 · 我们是怎么开发它的", level=1)

    _p(doc, "这类系统最大的难点不是“写出代码”，而是“知道代码到底错在哪、改了到底有没有变好”。"
            "我们没有凭感觉调参，而是建立了一套闭环工作流：让机器先跑、让人看清错误、把人的判断沉淀成"
            "“标准答案”、再用标准答案量化每次改动的好坏。下面是这套方法的六个环节。", bold=True)

    # 环节 1: 画廊
    doc.add_heading("环节一：先生成“画廊”，把 AI 看到的东西摊开给人看", level=2)
    _p(doc, "我们为每种能力（红绿灯、斑马线、车牌、跟踪）都做了“复核画廊”：一个本地网页，"
            "把 AI 在每一帧的判断（比如它认为这里是“绿灯”）连同原始画面截图并排铺出来。"
            "画廊是纯静态 HTML + 本地 JS，不依赖服务器，打开就能点。")
    _plain(doc, "就像把 AI 的“作业本”一页页摊在桌上：它在哪一帧觉得是绿灯、哪辆车它觉得停着，全都看得见，"
                "而不是只给一个最终分数。")
    _p(doc, "工程上，这由一个画廊基类（BaseGalleryBuilder）统一外壳，各能力子类"
            "（光灯/斑马线/车牌/跟踪画廊）只实现自己的标注逻辑。", size=10, color=(0x66, 0x66, 0x66))

    # 环节 2: 标注真值
    doc.add_heading("环节二：人工标注“真值”(标准答案)", level=2)
    _p(doc, "人在画廊里对着真实画面做判断——这一帧灯到底是红是绿、这辆车是不是违章车、车牌是多少——"
            "这些人工判断被收集起来，构成“真值”(Ground Truth, GT)。例如红绿灯真值是一份逐帧标注文件"
            "（覆盖全部视频、数百帧、数百个信号灯框，含颜色与“哪个灯管这条路”的标记）。")
    _plain(doc, "AI 要考试，总得先有“标准答案”。这一步就是人来出标准答案：对着画面一帧一帧告诉系统"
                "“这里正确答案是什么”。")

    # 环节 3: GT 体系(分层+可溯源)
    doc.add_heading("环节三：把真值做成“可溯源、不被污染”的体系", level=2)
    _p(doc, "真值不是一个随意的文件，而是一套分层体系：原始标注(source)、画廊修正(feedback)、"
            "迭代中发现的难例(badcases)分开存放；由一个带审计的构建脚本(build_gt.py)合并成"
            "权威真值(canonical)。合并时严守红线：输入只读绝不回写、每处改动可追溯到具体某一行反馈、"
            "合并前自动留快照和版本标签、对“绿↔红翻转”等高危改动默认不自动写入。")
    _plain(doc, "标准答案本身也要防出错、防被偷改。所以我们给它上了“档案管理”：谁改的、为什么改、"
                "改之前先存档，危险的改动必须人工签字——保证考试的标准答案永远可信。")
    _table(doc, ["层", "内容", "规则"],
           [
               ["source（原始）", "最初的人工标注", "只读"],
               ["feedback（修正）", "画廊里人工复核的更正", "只读，可追溯"],
               ["badcases（难例）", "迭代中发现的疑难帧", "累积复用"],
               ["canonical（权威）", "合并后的最终标准答案", "由脚本生成，带审计+快照"],
           ])

    # 环节 4: 对比错误帧
    doc.add_heading("环节四：拿 AI 的判断和真值逐帧对比，揪出“错误帧”", level=2)
    _p(doc, "有了标准答案，就能把 AI 的逐帧判断和真值对齐比对，自动找出它答错的帧——"
            "是把红灯看成了绿灯？还是把路边绿化带的绿色当成了绿灯？我们甚至专门做了“错误帧画廊”"
            "(mismatch gallery)，把 AI 答错的帧单独挑出来给人看，方便定位到底错在哪一类场景。")
    _plain(doc, "考完试，不只看得了几分，而是把每道错题都挑出来，看它为什么错——这样才知道该补哪里。")

    # 环节 5: 诊断根因 + 独立复现
    doc.add_heading("环节五：诊断根因，并要求“独立复现”才算数", level=2)
    _p(doc, "找到错误帧只是开始，还要诊断根因。例如我们发现某视频的“假绿”并不是红绿灯位置找错了，"
            "而是取色区域太大、把环境里的绿色也数了进去——这类结论必须用只读诊断脚本量化，"
            "而且由另一方“独立复现”到逐比特一致，才被采信。口头说“我改好了”不算数。")
    _plain(doc, "找病因不能靠猜。一个人说“我发现问题在这儿”，必须另一个人照着重做一遍、"
                "得出一模一样的数字，才算真的。这样杜绝了自说自话。")

    # 环节 6: AI 迭代 + 门禁
    doc.add_heading("环节六：AI 协同迭代，用“双重门禁”把关每一次改动", level=2)
    _p(doc, "这个项目由多个 AI 开发助手协同推进，各自在独立的工作副本(git worktree)里干活，互不干扰，"
            "改动都要署名。核心是“双重门禁”流程：先出方案过“方案门禁”(plan-gate)审查思路；"
            "实现后再过“效果门禁”(effect-gate)——必须在全部视频上真跑、指标不回退、"
            "尤其“误罚必须为 0”，由独立一方复现通过，才能合并进主干。")
    _plain(doc, "就像盖楼：先审图纸(方案门禁)，施工完再验收(效果门禁)，而且验收的人和施工的人不是同一个，"
                "必须亲自重新测一遍。任何会“冤枉好车”的改动，一票否决。")
    _p(doc, "正是这套流程，让系统在反复改动中始终守住“零误罚”这条底线。整个开发累计 300 多次提交，"
            "其中大量是诊断、复现、门禁记录，而非单纯写功能。", size=10, color=(0x66, 0x66, 0x66))

    # ---- 闭环图示(文字版) ----
    doc.add_heading("这套方法的闭环", level=2)
    _p(doc, "生成画廊 → 人工标注真值 → 对比揪出错误帧 → 诊断根因(独立复现) → "
            "AI 改进(双重门禁) → 回到画廊再验证……如此循环，每一圈系统都更准一点，而且每一步都看得见、"
            "可验证、可回溯。", bold=True)
    _plain(doc, "一句话：让机器先做、让人看懂错在哪、把人的判断变成标准答案、再用标准答案逼着机器一轮轮变好。")

    # ======================================================
    # 第三部分: 这套方法带来的结果
    # ======================================================
    doc.add_heading("第三部分 · 这套方法换来了什么", level=1)
    _p(doc, "在全部 11 段真实路口视频（含 2 段“无违章”负例）上的当前成绩：")
    _table(doc, ["口径", "指标", "含义"],
           [
               ["窗级（每次过街窗）", "F1 = 0.941，精度 = 1.000", "该报的基本都报，报的没有一个是错的"],
               ["车级（每辆违章车）", "召回 = 0.583，误罚 = 0", "多车场景能逐车识别，且绝不冤枉非违章车"],
               ["车牌", "命中 5/7", "读得清的都读对了，读不清的留空不瞎猜"],
           ])
    _p(doc, "更重要的是：这些数字都能被独立复现，每一个结论都能追溯到具体的诊断和门禁记录。"
            "系统“为什么这么判”始终说得清。")
    _plain(doc, "成绩不是重点，重点是这些成绩是“可信的、可复查的”——这正是严肃工程和“碰运气调出来”的根本区别。")

    doc.add_heading("给老师与评委的话", level=1)
    _p(doc, "这个项目真正的亮点，不止是“做了一个能抓违章的程序”，而是展示了一套严谨的工程方法论："
            "把复杂问题拆成可验证的小步、用人工真值当标尺、让每次改动都经过独立复现与门禁、"
            "并始终守住“零误罚”的伦理底线。")
    _p(doc, "对学生而言，这说明了人工智能不是“魔法黑箱”：它的每一步判断都可以被看见、被质疑、被改进。"
            "学会用“数据+验证+迭代”的方式解决问题，比做出某一个具体程序更有价值。")

    doc.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))), "docs", "manuals"))
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    out = build(os.path.join(args.out, "开发过程文档.docx"))
    print("[生成]", out)


if __name__ == "__main__":
    main()
